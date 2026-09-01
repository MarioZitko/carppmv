#!/usr/bin/env python
"""CLI: ingest one xlsx catalogue file into the Catalogue database table.

Usage:
    uv run python scripts/ingest_catalogue.py path/to/catalogue.xlsx
    uv run python scripts/ingest_catalogue.py path/to/catalogue.xlsx --co2-standard NEDC

OPENROUTER_API_KEY must be set in the environment (or .env) before running.

Design notes:
- One LLM call per sheet (map_sheet_columns), never per row — bulk data never
  sent to the model, keeping cost near zero.
- Sheet-level confidence gate: 0.6. Below this, the LLM could not reliably
  identify the required fields; ingesting would produce garbage even with every
  row flagged needs_review=True. Sheets in [0.6, 0.7) are ingested but all
  rows are flagged needs_review=True (same threshold as canonical_schema's
  NEEDS_REVIEW_CONFIDENCE_THRESHOLD). Below 0.6 the sheet is skipped.
- co2_standard is not derivable from the Excel rows — must be supplied per-file
  via --co2-standard (default WLTP; WLTP mandatory for EU type-approvals from 2017).
- Duplicate rows (same brand/model/variant/valid_from) are silently skipped via
  ON CONFLICT DO NOTHING rather than failing the whole sheet.
"""

import argparse
import asyncio
import sys
from collections import Counter
from pathlib import Path

import openpyxl
from sqlalchemy.dialects.postgresql import insert as pg_insert

# Ensure project root is on sys.path so `app.*` imports work when run directly.
sys.path.insert(0, str(Path(__file__).parent.parent))

from app.catalogue.canonical_schema import CanonicalRow, FuelCategory, apply_mapping
from app.catalogue.llm_mapper import map_sheet_columns
from app.catalogue.matching import build_match_key
from app.core.config import get_settings
from app.db.models import Catalogue, CO2Standard, FuelType
from app.db.session import AsyncSessionLocal

# Sheet-level confidence gate — see module docstring for rationale.
SHEET_CONFIDENCE_THRESHOLD = 0.6

_FUEL_CATEGORY_TO_DB: dict[FuelCategory, FuelType | None] = {
    FuelCategory.DIESEL: FuelType.DIESEL,
    FuelCategory.PETROL: FuelType.PETROL,
    FuelCategory.PETROL_HYBRID: FuelType.PETROL,          # mild hybrid taxed as petrol
    FuelCategory.PETROL_PLUG_IN_HYBRID: FuelType.PETROL,  # plug-in taxed as petrol
    FuelCategory.ELECTRIC: None,   # PPMV-exempt — no Catalogue row needed
    FuelCategory.UNKNOWN: None,    # needs_review=True already set; don't insert bad data
}


def _to_catalogue_dict(row: CanonicalRow, co2_standard: CO2Standard) -> dict | None:
    """Map a CanonicalRow to a dict ready for a Catalogue INSERT.

    Returns None for ELECTRIC and UNKNOWN rows, which should not be inserted.
    PETROL_HYBRID and PETROL_PLUG_IN_HYBRID both become FuelType.PETROL — they
    are taxed as petrol in the PPMV engine (see FuelCategory docstring).

    model / variant come from model_name / full_name with type_code as fallback:
    Catalogue's unique key is (brand, model, variant, valid_from), so using
    full_name as variant gives the best per-variant granularity for lookup.

    match_key is the normalized brand+model+variant used for listing matching
    (see app/catalogue/matching.py); it's built from the SAME model/variant
    strings stored below, so a lookup-time key reconstructed from those columns
    matches it exactly. power_kw is carried through as a matching disambiguator.
    """
    fuel_type = _FUEL_CATEGORY_TO_DB.get(row.fuel_category)
    if fuel_type is None:
        return None

    model = row.model_name or row.type_code or "unknown"
    variant = row.full_name or row.type_code or row.model_name or "unknown"

    return {
        "brand": row.brand,
        "model": model,
        "variant": variant,
        "match_key": build_match_key(row.brand, model, variant),
        "price_eur": row.price_eur,
        "co2_g_km": row.co2_g_km,
        "co2_standard": co2_standard,
        "fuel_type": fuel_type,
        "power_kw": row.power_kw,
        "valid_from": row.valid_from,
        "source_file": row.source_file,
        "source_currency": row.price_source_currency,
    }


async def _insert_rows(rows: list[dict]) -> int:
    """Bulk-insert with ON CONFLICT DO NOTHING. Returns number of rows inserted."""
    if not rows:
        return 0
    async with AsyncSessionLocal() as session:
        stmt = pg_insert(Catalogue).values(rows).on_conflict_do_nothing(
            constraint="uq_catalogue_lookup_key"
        )
        result = await session.execute(stmt)
        await session.commit()
        return result.rowcount if result.rowcount is not None else 0


async def ingest_file(xlsx_path: Path, co2_standard: CO2Standard) -> None:
    settings = get_settings()
    if not settings.openrouter_api_key:
        print(
            "ERROR: OPENROUTER_API_KEY is not set.\n"
            "Add it to .env or export it before running.",
            file=sys.stderr,
        )
        sys.exit(1)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheet_names = wb.sheetnames
    print(f"Opened: {xlsx_path.name} — {len(sheet_names)} sheet(s): {sheet_names}")
    print(f"CO2 standard: {co2_standard.value}")
    print()

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        all_rows = list(ws.iter_rows(values_only=True))

        if not all_rows:
            print(f"[{sheet_name}] SKIP — empty sheet")
            continue

        header_row = [str(c) if c is not None else "" for c in all_rows[0]]
        data_rows = all_rows[1:]

        if not data_rows:
            print(f"[{sheet_name}] SKIP — no data rows")
            continue

        print(f"[{sheet_name}] Mapping {len(header_row)} columns via LLM...", end=" ", flush=True)
        try:
            mapping = await map_sheet_columns(header_row, data_rows[:5])
        except Exception as exc:
            print(f"FAILED\n  LLM error: {exc}")
            continue

        print(f"confidence={mapping.confidence:.2f}")

        if mapping.confidence < SHEET_CONFIDENCE_THRESHOLD:
            print(
                f"[{sheet_name}] SKIP — confidence {mapping.confidence:.2f} "
                f"< threshold {SHEET_CONFIDENCE_THRESHOLD:.2f}\n"
                f"  Notes: {mapping.notes}"
            )
            continue

        skip_log: list[tuple[int, str]] = []
        canonical_rows = apply_mapping(
            rows=data_rows,
            header_row=header_row,
            mapping=mapping,
            source_file=str(xlsx_path),
            source_sheet=sheet_name,
            skip_log=skip_log,
        )

        rows_read = len(data_rows)
        rows_kept = len(canonical_rows)
        rows_skipped = len(skip_log)
        skip_counts: Counter[str] = Counter(reason for _, reason in skip_log)

        catalogue_dicts: list[dict] = []
        not_insertable = 0
        for row in canonical_rows:
            d = _to_catalogue_dict(row, co2_standard)
            if d is None:
                not_insertable += 1
            else:
                catalogue_dicts.append(d)

        try:
            inserted = await _insert_rows(catalogue_dicts)
            duplicates = len(catalogue_dicts) - inserted
        except Exception as exc:
            # SQLAlchemy exceptions stringify to the full SQL statement plus every
            # bound parameter — unreadable for a batch of hundreds of rows.
            orig = getattr(exc, "orig", None)
            short = f"{type(orig).__name__}: {orig}" if orig is not None else f"{type(exc).__name__}: {exc}"
            print(f"[{sheet_name}] DB insert failed: {short}")
            inserted = 0
            duplicates = 0

        print(
            f"[{sheet_name}] "
            f"read={rows_read}  kept={rows_kept}  "
            f"inserted={inserted}  duplicates_skipped={duplicates}  "
            f"not_insertable(electric/unknown)={not_insertable}  "
            f"dropped={rows_skipped}"
        )
        for reason, count in skip_counts.most_common():
            print(f"  dropped[{reason}]={count}")
        if mapping.notes:
            print(f"  llm_notes: {mapping.notes}")
        print()


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Ingest a brand Excel catalogue into the Catalogue DB table.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    parser.add_argument("xlsx_path", type=Path, help="Path to the .xlsx file to ingest.")
    parser.add_argument(
        "--co2-standard",
        choices=["NEDC", "WLTP"],
        default="WLTP",
        metavar="{NEDC,WLTP}",
        help=(
            "CO2 measurement standard for all sheets in this file (default: WLTP). "
            "Cannot be inferred from the Excel — supply it based on when the catalogue "
            "was produced (WLTP mandatory for EU type-approvals from September 2017)."
        ),
    )
    args = parser.parse_args()

    if not args.xlsx_path.exists():
        print(f"ERROR: file not found: {args.xlsx_path}", file=sys.stderr)
        sys.exit(1)

    asyncio.run(ingest_file(args.xlsx_path, CO2Standard[args.co2_standard]))


if __name__ == "__main__":
    main()
